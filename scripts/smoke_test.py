"""End-to-end smoke test for the core pipeline (no Ollama required).

Creates a temp folder of varied files, indexes them with AI tagging disabled,
and runs keyword searches — exercising the scanner, every content extractor,
the SQLite/FTS5 index, and the search engine. Run:

    .venv\\Scripts\\python.exe scripts\\smoke_test.py

Exits non-zero if any expected match is missing.
"""

from __future__ import annotations

import sys
import tempfile
from pathlib import Path

from filesearch.config import Config
from filesearch.indexer import Indexer
from filesearch.search import SearchEngine


def make_sample_files(root: Path) -> None:
    (root / "notes.txt").write_text(
        "Meeting notes about the Q3 budget and the marketing strategy for next year.",
        encoding="utf-8",
    )
    (root / "readme.md").write_text(
        "# Project Phoenix\n\nA tool for analyzing customer churn and retention.",
        encoding="utf-8",
    )
    (root / "app.py").write_text(
        "def calculate_tax(amount):\n    return amount * 0.2\n\n# invoice processing logic\n",
        encoding="utf-8",
    )
    (root / "data.json").write_text(
        '{"customer": "Acme", "invoice_total": 1200, "currency": "USD"}',
        encoding="utf-8",
    )
    (root / "sales.csv").write_text(
        "month,revenue\nJan,1000\nFeb,1500\nMar,1750\n", encoding="utf-8"
    )
    # docx
    import docx

    d = docx.Document()
    d.add_paragraph("Annual financial report. Revenue grew 20% year over year.")
    d.save(str(root / "report.docx"))
    # pdf
    import fitz

    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "This agreement is between Acme Corp and Beta LLC.")
    doc.save(str(root / "contract.pdf"))
    doc.close()
    # xlsx
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["product", "units"])
    ws.append(["widget", 42])
    wb.save(str(root / "inventory.xlsx"))
    # image (metadata-only without a vision model)
    from PIL import Image

    Image.new("RGB", (64, 64), (120, 160, 200)).save(str(root / "chart.png"))
    # a hidden file + a hidden dir that must be EXCLUDED
    (root / ".secret.txt").write_text("api_key=DO_NOT_INDEX", encoding="utf-8")
    hidden = root / ".hidden_dir"
    hidden.mkdir()
    (hidden / "inside.txt").write_text("should not be indexed", encoding="utf-8")


def main() -> int:
    tmp = Path(tempfile.mkdtemp(prefix="filesearch_smoke_"))
    root = tmp / "docs"
    root.mkdir()
    make_sample_files(root)

    cfg = Config(
        scan_roots=[str(root)],
        db_path=str(tmp / "index.db"),
        enable_ai_tagging=False,
        enable_embeddings=False,
    )
    indexer = Indexer(cfg)
    stats = indexer.run_scan()
    print(f"scan stats: {stats['scanned']} scanned, {stats['skipped']} skipped, {stats['errors']} errors")

    engine = SearchEngine(cfg)
    checks = {
        "budget": "notes.txt",
        "phoenix": "readme.md",
        "tax": "app.py",
        "invoice": "app.py",
        "revenue": "sales.csv",
        "agreement": "contract.pdf",
        "widget": "inventory.xlsx",
    }
    failures = []
    for term, expected in checks.items():
        hits = engine.search(term, mode="keyword", limit=10)
        names = [h.name for h in hits]
        ok = expected in names
        print(f"  query '{term}': {'OK' if ok else 'MISSING'} -> {names}")
        if not ok:
            failures.append((term, expected, names))

    # The hidden file and hidden dir must NOT be indexed.
    all_paths = engine.index.all_paths()
    leaked = [p for p in all_paths if ".secret.txt" in p or ".hidden_dir" in p]
    if leaked:
        failures.append(("hidden-exclusion", "none", leaked))
        print(f"  EXCLUSION FAIL: hidden paths leaked into index: {leaked}")
    else:
        print("  hidden-file/dir exclusion: OK")

    expected_count = 9  # the 9 visible files (2 hidden entries excluded)
    print(f"  indexed file count: {stats['scanned']} (expected {expected_count})")

    engine.index.close()
    indexer.index.close()
    import shutil

    shutil.rmtree(tmp, ignore_errors=True)

    if failures:
        print(f"\nSMOKE TEST FAILED: {len(failures)} problem(s)")
        return 1
    print("\nSMOKE TEST PASSED")
    return 0


if __name__ == "__main__":
    sys.exit(main())
